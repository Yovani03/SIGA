from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django.utils import timezone
from .models import Viaje, ConfiguracionBono, Bitacora
from .serializers import ViajeSerializer, ConfiguracionBonoSerializer, BitacoraSerializer
from operadores.models import Operador
from vehiculos.models import UnidadTractocamion
import openpyxl
import os
import datetime
from django.core.files import File
from django.conf import settings
from tempfile import NamedTemporaryFile

class ViajeViewSet(viewsets.ModelViewSet):
    queryset = Viaje.objects.all()
    serializer_class = ViajeSerializer

    def perform_create(self, serializer):
        # Si no se proporciona fecha_salida, usar la actual
        fecha_salida = self.request.data.get('fecha_salida')
        if not fecha_salida:
            viaje = serializer.save(fecha_salida=timezone.now())
        else:
            viaje = serializer.save()
            
        # Actualizar estatus del operador
        operador = viaje.operador
        operador.estatus = 'viaje'
        operador.save()
        
        # Actualizar estatus del ayudante si existe
        if viaje.ayudante:
            ayudante = viaje.ayudante
            ayudante.estatus = 'viaje'
            ayudante.save()

    @action(detail=True, methods=['post'])
    def avanzar_estatus(self, request, pk=None):
        viaje = self.get_object()
        if viaje.completado:
            return Response({'error': 'Este viaje ya ha sido completado.'}, status=status.HTTP_400_BAD_REQUEST)
        
        # En el nuevo flujo, cualquier avance desde 'en_ruta' significa completar el viaje (llegada a CEDIS)
        viaje.estatus = 'completado'
        viaje.completado = True
        
        # Si se envía fecha_llegada en el body, usarla. Si no, usar la actual.
        fecha_llegada = request.data.get('fecha_llegada')
        if fecha_llegada:
            viaje.fecha_llegada = fecha_llegada
        else:
            viaje.fecha_llegada = timezone.now()
        
        viaje.save()

        # Liberar personal
        operador = viaje.operador
        operador.estatus = 'patio'
        operador.save()
        
        if viaje.ayudante:
            ayudante = viaje.ayudante
            ayudante.estatus = 'patio'
            ayudante.save()
            
        return Response(ViajeSerializer(viaje).data)

    @action(detail=True, methods=['post'])
    def set_ayudante(self, request, pk=None):
        viaje = self.get_object()
        ayudante_id = request.data.get('ayudante_id')
        
        if not ayudante_id:
            return Response({'error': 'ID de ayudante requerido.'}, status=status.HTTP_400_BAD_REQUEST)
            
        try:
            ayudante = Operador.objects.get(id=ayudante_id)
            if ayudante.estatus == 'viaje':
                return Response({'error': 'El ayudante ya est en un viaje.'}, status=status.HTTP_400_BAD_REQUEST)
                
            viaje.ayudante = ayudante
            viaje.save()
            
            # Marcar ayudante como ocupado
            ayudante.estatus = 'viaje'
            ayudante.save()
            
            return Response(ViajeSerializer(viaje).data)
        except Operador.DoesNotExist:
            return Response({'error': 'Operador no encontrado.'}, status=status.HTTP_404_NOT_FOUND)

    @action(detail=True, methods=['post'])
    def registrar_llegada(self, request, pk=None):
        viaje = self.get_object()
        viaje.estatus = 'completado'
        viaje.completado = True
        
        fecha_llegada = request.data.get('fecha_llegada')
        if fecha_llegada:
            viaje.fecha_llegada = fecha_llegada
        else:
            viaje.fecha_llegada = timezone.now()
            
        viaje.save()

        operador = viaje.operador
        operador.estatus = 'patio'
        operador.save()
        
        if viaje.ayudante:
            ayudante = viaje.ayudante
            ayudante.estatus = 'patio'
            ayudante.save()

        return Response(ViajeSerializer(viaje).data)

class ConfiguracionBonoViewSet(viewsets.ModelViewSet):
    queryset = ConfiguracionBono.objects.all()
    serializer_class = ConfiguracionBonoSerializer
    ordering = ['capacidad']

class BitacoraViewSet(viewsets.ModelViewSet):
    queryset = Bitacora.objects.all()
    serializer_class = BitacoraSerializer

    @action(detail=True, methods=['get'])
    def pdf(self, request, pk=None):
        bitacora = self.get_object()
        
        context = {
            'numero_economico': bitacora.vehiculo.numero_economico,
            'placas': bitacora.vehiculo.placas or 'Sin placas',
            'folio': bitacora.folio,
            'semana': f"{bitacora.fecha_inicio.strftime('%d/%m/%Y')} al {bitacora.fecha_fin.strftime('%d/%m/%Y')}"
        }
        
        from django.template.loader import render_to_string
        from django.http import HttpResponse
        import xhtml2pdf.pisa as pisa
        import io
        
        html_string = render_to_string('bitacora_base.html', context)
        
        result = io.BytesIO()
        pdf = pisa.pisaDocument(io.BytesIO(html_string.encode("UTF-8")), result)
        
        if not pdf.err:
            response = HttpResponse(result.getvalue(), content_type='application/pdf')
            response['Content-Disposition'] = f'inline; filename="Bitacora_{bitacora.vehiculo.numero_economico}_Folio{bitacora.folio}.pdf"'
            return response
        return Response({'error': 'Error generating PDF'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=True, methods=['get'])
    def descargar_copia(self, request, pk=None):
        bitacora = self.get_object()
        if not bitacora.archivo:
            return Response({'error': 'La bitácora no tiene archivo.'}, status=status.HTTP_404_NOT_FOUND)
            
        bitacora.copias_descargadas += 1
        bitacora.save()
        
        try:
            wb = openpyxl.load_workbook(bitacora.archivo.path)
            sheet = wb.active
            
            from openpyxl.drawing.image import Image as OpenpyxlImage
            from PIL import Image, ImageDraw, ImageFont
            import os
            from tempfile import NamedTemporaryFile
            
            # Generar imagen de marca de agua transparente
            img = Image.new('RGBA', (800, 600), (255, 255, 255, 0))
            try:
                font = ImageFont.truetype("arial.ttf", 90)
            except IOError:
                font = ImageFont.load_default()
                
            txt = Image.new('RGBA', (600, 200), (255, 255, 255, 0))
            d = ImageDraw.Draw(txt)
            # Negro semitransparente
            d.text((10, 10), f"COPIA {bitacora.copias_descargadas}", font=font, fill=(0, 0, 0, 90))
            
            txt = txt.rotate(35, expand=1)
            img.paste(txt, (100, 150), txt)
            
            with NamedTemporaryFile(delete=False, suffix='.png') as tmp_img:
                img.save(tmp_img.name)
                tmp_img_path = tmp_img.name
                
            xl_img = OpenpyxlImage(tmp_img_path)
            sheet.add_image(xl_img, 'D10')
            
            from django.http import HttpResponse
            import tempfile
            
            with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_out:
                wb.save(tmp_out.name)
                tmp_out_path = tmp_out.name
                
            with open(tmp_out_path, 'rb') as f:
                response = HttpResponse(f.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                response['Content-Disposition'] = f'attachment; filename="Copia_{bitacora.copias_descargadas}_Bitacora_U{bitacora.vehiculo.numero_economico}_Folio{bitacora.folio}.xlsx"'
            
            try:
                os.unlink(tmp_img_path)
                os.unlink(tmp_out_path)
            except:
                pass
                
            return response
            
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


    @action(detail=False, methods=['post'])
    def generar(self, request):
        vehiculo_id = request.data.get('vehiculo_id')
        fecha_str = request.data.get('fecha_inicio')

        if not vehiculo_id or not fecha_str:
            return Response({'error': 'vehiculo_id y fecha_inicio son requeridos.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            vehiculo = UnidadTractocamion.objects.get(id=vehiculo_id)
        except UnidadTractocamion.DoesNotExist:
            return Response({'error': 'Vehículo no encontrado.'}, status=status.HTTP_404_NOT_FOUND)

        try:
            fecha_inicio = datetime.datetime.strptime(fecha_str, '%Y-%m-%d').date()
        except ValueError:
            return Response({'error': 'Formato de fecha inválido. Usa YYYY-MM-DD.'}, status=status.HTTP_400_BAD_REQUEST)

        # Ensure fecha_inicio is a Friday
        if fecha_inicio.weekday() != 4: # 4 is Friday
            # Adjust to the most recent Friday or next Friday?
            # Normally the frontend sends the correct Friday. If not, error out or auto-correct.
            return Response({'error': 'La fecha de inicio debe ser un viernes.'}, status=status.HTTP_400_BAD_REQUEST)

        fecha_fin = fecha_inicio + datetime.timedelta(days=6)

        # Generate folio
        last_bitacora = Bitacora.objects.filter(vehiculo=vehiculo).order_by('-folio').first()
        folio = 1 if not last_bitacora else last_bitacora.folio + 1

        # Check if already exists for this date and vehicle
        existing = Bitacora.objects.filter(vehiculo=vehiculo, fecha_inicio=fecha_inicio).first()
        if existing:
            return Response(BitacoraSerializer(existing).data)

        # Excel Manipulation
        template_path = settings.BASE_DIR / 'logistica' / 'plantillas' / 'FORMATO CEDIS Y TRANSPORTES FINAL.xlsx'
        if not template_path.exists():
            return Response({'error': f'Template de Excel no encontrado: {template_path}'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        wb = openpyxl.load_workbook(str(template_path))
        sheet = wb.active
        sheet.title = vehiculo.numero_economico or "Bitacora"

        meses = ['ene', 'feb', 'mar', 'abr', 'may', 'jun', 'jul', 'ago', 'sep', 'oct', 'nov', 'dic']
        str_fecha = f"{fecha_inicio.day} de {meses[fecha_inicio.month - 1]} al {fecha_fin.day} de {meses[fecha_fin.month - 1]}"

        # Write to the merged cells directly (top-left cell of each merge)
        sheet['A1'] = f"Nº ECONOMICO: {vehiculo.numero_economico}"
        sheet['A2'] = f"PLACAS: {vehiculo.placas or 'S/N'}"
        sheet['F1'] = "AUTOTRANSPORTES COLUMBIA"
        sheet['L1'] = f"FOLIO: {folio:03d}"
        sheet['L2'] = f"SEMANA: {str_fecha}"

        # Save to temp file
        with NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp:
            wb.save(tmp.name)
            tmp_path = tmp.name

        # Create record
        bitacora = Bitacora(
            vehiculo=vehiculo,
            folio=folio,
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_fin
        )
        
        with open(tmp_path, 'rb') as f:
            file_name = f"Bitacora_U{vehiculo.numero_economico}_Folio{folio}_{fecha_inicio.strftime('%Y%m%d')}.xlsx"
            bitacora.archivo.save(file_name, File(f))
            
        bitacora.save()
        os.unlink(tmp_path)

        return Response(BitacoraSerializer(bitacora).data, status=status.HTTP_201_CREATED)
